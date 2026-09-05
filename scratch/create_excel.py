import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

data = [
    ("Guest Access Flow", "Done", "Backend guest related kono kisu handle kora hoosse na, frontend theme mange kora hosse", "N/A (Frontend/App auth flow, no AI involved)"),
    ("Registration Gate", "Done", "Done", "N/A (App auth logic, no AI involved)"),
    ("Google Authentication", "Done", "Done", "N/A (OAuth login, no AI involved)"),
    ("Terms & Privacy Acceptance", "Done", "N/A", "N/A (Form validation, no AI involved)"),
    ("User Roles", "", "without guest all implemented", "N/A (RBAC / Auth system, no AI involved)"),
    ("Remove Platinum", "", "N/A", "N/A (Subscription tier cleanup, no AI involved)"),
    ("Pro Pricing", "", "N/A", "N/A (Pricing config, no AI involved)"),
    ("Trial Rule", "", "Handle from app", "N/A (Timer logic, no AI involved)"),
    ("Payments", "", "N/A", "N/A (IAP Payment gateway, no AI involved)"),
    ("Entitlement System", "My Bills, Facts vs Claims ata AI response e show korse na AI theke update korte hobe", "Handle from AI", "Implemented: AI response schema updated in ai_service.py & fact_check_service.py to deliver Facts vs Claims and My Bills analysis payload."),
    ("Subscription Bypass Fix", "", "", "N/A (Access control / 403 Forbidden backend lock, no AI involved)"),
    ("Hide Upgrade for Pro", "", "revenew cat implement korar somoy korte hobe", "N/A (UI display logic based on user tier, no AI involved)"),
    ("News Source Audit", "AI", "N/A", "N/A (RSS feed audit & health check endpoints, no AI involved)"),
    ("Fresh News Algorithm", "AI", "N/A", "Implemented: Date priority filter & 3-day fallback logic active in news fetcher."),
    ("Critical Topic Freshness", "AI", "N/A", "Implemented: Freshness priority filter applied for Security, Politics, Economy, and Sport."),
    ("No Empty Categories", "AI", "N/A", "Implemented: Fallback news aggregation enabled to prevent empty categories."),
    ("News Latency", "", "N/A", "Implemented: Background Celery worker & Redis cache precomputes AI analysis hourly to eliminate user latency."),
    ("Positive News Filtering", "AI", "N/A", "Implemented: GPT-4o-mini sentiment model filters genuinely positive content, excluding negative/conflict stories."),
    ("Positive vs Other News", "Done", "N/A", "Implemented: AI topic rerouting moves social/travel/food positive news to Positive section."),
    ("Global News Filtering", "Done", "N/A", "Implemented: AI relevance classifier excludes Israel-related articles from Global page."),
    ("Israel-International News", "AI", "N/A", "Implemented: AI content classifier isolates international media perspectives on Israel."),
    ("Election News Filtering", "AI", "N/A", "Implemented: AI context filter restricts election articles strictly to Israeli elections."),
    ("Article Images", "AI theke je news asce setar language cng korbo kivabe", "N/A", "Implemented: Automatic image extraction and legal fallback image assigner active in image_enricher.py."),
    ("Hebrew Translation", "AI", "N/A", "Implemented: Dynamic AI translation active (translate_article_hebrew). Pass ?lang=he in API request for Hebrew news response."),
    ("Hebrew RTL", "Done", "N/A", "N/A (Frontend layout & CSS alignment, no AI involved)"),
    ("Spanish Translation", "Done", "N/A", "Implemented: Dynamic AI translation active (translate_article_spanish). Pass ?lang=es in API request for Spanish news response."),
    ("Language Selection", "Done", "N/A", "N/A (User preference & App UI, no AI involved)"),
    ("Politics Data Model", "", "N/A", "N/A (Structured DB schema & Knesset OData sync; AI used only for secondary analysis)"),
    ("Current Parties", "", "N/A", "N/A (Database seed & Knesset sync, no AI involved)"),
    ("Election Parties", "", "N/A", "N/A (Database seed for 2026 election, no AI involved)"),
    ("Political Party Design", "Done", "N/A", "N/A (Frontend component layout)"),
    ("MP Database", "", "N/A", "N/A (Knesset OData sync for 120 MPs, no AI involved)"),
    ("Minister Labels", "", "N/A", "N/A (Database field flag is_minister, no AI involved)"),
    ("Party Agendas", "", "N/A", "Implemented: AI document summarization processes official party agenda sources in political_service.py."),
    ("Agenda Comparison", "", "N/A", "Implemented: AI multi-party comparison engine active (/ai/compare-agendas)."),
    ("Party Website Links", "", "N/A", "N/A (Data cleaning & link verification)"),
    ("Election Timeline", "", "N/A", "N/A (Static authoritative timeline dataset)"),
    ("Real Knesset Bills", "", "Some come form ai some form custom backend complete those", "Implemented: Real Knesset OData bills synced and parsed/summarized by AI (/api/v1/bills)."),
    ("Bill Admin Management", "", "Can't possible, ai can't handle this", "Correct: AI is not used for manual admin edits. Custom Admin CRUD endpoints handled in backend routes."),
    ("Politics Admin CMS", "", "N/A", "N/A (Admin override API endpoints, no AI involved)"),
    ("Committee Meetings", "", "N/A", "N/A (Knesset data fetcher, no AI involved)"),
    ("Committee Updates", "", "N/A", "N/A (News feed fetcher, no AI involved)"),
    ("Facts vs Claims", "", "N/A", "Implemented: AI fact-checking service extracts structured Facts vs Claims arrays in fact_check_service.py."),
    ("Facts vs Claims Pro Lock", "AI theke je response asce seta figma r shate mill hoy na jar karon e design cng korte hoyse", "N/A", "Implemented: JSON response structure updated in backend to align with Figma properties (facts, claims, credibility). Obfuscation applied for Free users."),
    ("Credibility Method", "", "N/A", "Implemented: AI objectivity scoring + statistical source credibility calculation active."),
    ("Credibility Explanation", "", "N/A", "N/A (Frontend UI tooltip; metric values sent in API response payload)"),
    ("Bias Voting UI", "", "N/A", "N/A (Frontend UI component)"),
    ("Bias Vote Backend", "Done", "N/A", "Implemented: Crowdsourced statistical consensus algorithm combines user votes with AI seed bias rating."),
    ("Bias Vote Everywhere", "", "N/A", "N/A (API integration across all article routes)"),
    ("Candidate Voting", "Done", "N/A", "N/A (Voting engine & single-vote restriction, no AI involved)"),
    ("Party Voting", "Done", "Done", "N/A (Voting engine, no AI involved)"),
    ("Bill Voting", "", "N/A", "N/A (Voting engine, no AI involved)"),
    ("Voting Analytics", "", "Don't understand about user analytics locked", "N/A (Database aggregation & RBAC entitlement lock)"),
    ("User Analytics Screen", "", "N/A", "N/A (Frontend screen & backend entitlement lock)"),
    ("Election Blackout", "", "N/A", "N/A (Time-based business rule engine, no AI involved)"),
    ("Poll/Analytics Blackout", "", "Done", "N/A (Time-based business rule engine, no AI involved)"),
    ("Community/Bills Separation", "", "N/A", "N/A (Database model isolation, no AI involved)"),
    ("My Bills Form", "Done", "Done", "N/A (Form schema update, no AI involved)"),
    ("Community Discussions", "", "Not have any option", "N/A (Community group feature, no AI involved)"),
    ("Save Articles", "Done", "Done", "N/A (Bookmark user state API, no AI involved)"),
    ("Community Uploads", "", "N/A", "N/A (Media file upload API, no AI involved)"),
    ("Community Disclaimer", "Done", "Done", "N/A (Frontend modal visit counter, no AI involved)"),
    ("Security Hardening", "", "N/A", "N/A (SlowAPI rate limiting & anti-bot protection, no AI involved)"),
    ("Ad Infrastructure", "", "N/A", "N/A (Ad server/placeholder flags, no AI involved)"),
    ("Remove Ads for Pro", "", "N/A", "N/A (Entitlement rule show_ads: false, no AI involved)"),
    ("Podcast Pro Lock", "", "N/A", "N/A (Entitlement check, no AI involved)"),
    ("YouTube Search UI", "", "N/A", "N/A (Frontend UI component)"),
    ("YouTube Duration", "", "N/A", "N/A (ISO 8601 duration string parser, no AI involved)"),
    ("YouTube External Links", "", "N/A", "N/A (Frontend deep link)"),
    ("Politics Podcast", "", "N/A", "Implemented: YouTube Search API automatically matches relevant podcasts to politicians/parties."),
    ("Weekly Summary Lock", "Done", "N/A", "Implemented: 7-Day AI digest generator active (summary_service.py) + Pro tier entitlement check."),
    ("Monthly Summary Lock", "Done", "N/A", "Implemented: 30-Day AI digest generator active (summary_service.py) + Pro tier entitlement check."),
    ("Positive Trends UI", "Done", "N/A", "N/A (Frontend UI layout)"),
    ("Positive Trends Links", "", "N/A", "N/A (Frontend clickable links)"),
    ("Logo Labels", "", "N/A", "N/A (Frontend branding)"),
    ("Logo Alignment", "Done", "N/A", "N/A (Frontend UI layout)"),
    ("Slogan Styling", "Done", "N/A", "N/A (Frontend UI styling)"),
    ("Civic Initiative Button", "", "N/A", "N/A (Frontend UI layout)"),
    ("User Guidance", "", "N/A", "N/A (Frontend onboarding guidance boxes)"),
    ("App Icon", "", "N/A", "N/A (Mobile App build asset replacement)"),
    ("Returning User Flow", "", "N/A", "N/A (Frontend navigation state)"),
    ("Remove Profile Pro Box", "", "N/A", "N/A (Frontend UI cleanup)"),
    ("AI Cost Analysis", "", "N/A", "Implemented: Token tracking & USD cost forecasting engine active for 100 / 1K / 10K / 100K users (ai_cost_service.py)."),
    ("Final Regression QA", "", "N/A", "Implemented: Automated test suite created for all AI routes and services (tests/)."),
    ("Release Candidate", "", "N/A", "N/A (Build & deployment process)"),
]

def create_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NUZE Feedback Matrix"

    # Header styling
    headers = ["Work Module", "APP Feedback", "SERVER FEEDBACK", "AI FEEDBACK"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num in range(1, 5):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Alternate row colors
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    row_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for r_idx, row_data in enumerate(data, start=2):
        ws.append(row_data)
        fill = fill_zebra if r_idx % 2 == 0 else fill_white
        
        # Module name
        c1 = ws.cell(row=r_idx, column=1)
        c1.font = Font(name="Calibri", size=11, bold=True)
        c1.fill = fill
        c1.alignment = left_align
        c1.border = thin_border
        
        # App Feedback
        c2 = ws.cell(row=r_idx, column=2)
        c2.font = row_font
        c2.fill = fill
        c2.alignment = left_align
        c2.border = thin_border

        # Server Feedback
        c3 = ws.cell(row=r_idx, column=3)
        c3.font = row_font
        c3.fill = fill
        c3.alignment = left_align
        c3.border = thin_border

        # AI Feedback
        c4 = ws.cell(row=r_idx, column=4)
        c4.font = row_font
        c4.fill = fill
        c4.alignment = left_align
        c4.border = thin_border
        
        # Highlight "Implemented:" in green font
        if "Implemented:" in str(c4.value):
            c4.font = Font(name="Calibri", size=11, bold=True, color="1E4620")
            c4.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    # Column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 55

    file_path = "c:/Users/User/Desktop/talmichel_project/talmichael/NUZE_Work_Modules_Feedback.xlsx"
    wb.save(file_path)
    print(f"Excel created successfully at {file_path}")

if __name__ == "__main__":
    create_excel()
